"""Текст DOCX: абзацы и ячейки таблиц. Не LibreOffice и не PDF.

Текстовые рамки, content control и колонтитулы python-docx в основной
текст не отдаёт. Их наличие записывается, текст не выдумывается.
"""

from __future__ import annotations

import zipfile
from dataclasses import dataclass
from io import BytesIO

from docx import Document
from docx.opc.exceptions import PackageNotFoundError

from kontur.application.extractors.number import PageToken
from kontur.domain.models import ExtractionEngine

_UNIT = (
    (0.05, 0.05),
    (0.45, 0.05),
    (0.45, 0.08),
    (0.05, 0.08),
)


class DocxReadError(ValueError):
    """Файл с расширением docx не разобран. Это не неподдерживаемый формат."""


@dataclass(frozen=True, slots=True)
class DocxText:
    tokens: tuple[PageToken, ...]
    skipped: tuple[str, ...]


def read_docx_bytes(raw: bytes) -> DocxText:
    """Слова абзацев и таблиц. Пустой текст — пустые tokens, не исключение."""

    try:
        document = Document(BytesIO(raw))
    except (PackageNotFoundError, zipfile.BadZipFile, ValueError, OSError) as exc:
        raise DocxReadError(f"DOCX: не прочитан ({exc.__class__.__name__})") from exc
    words: list[str] = []
    for paragraph in document.paragraphs:
        words.extend(_words(paragraph.text))
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    words.extend(_words(paragraph.text))
    tokens = tuple(_token(word, index) for index, word in enumerate(words))
    return DocxText(tokens=tokens, skipped=_skipped_parts(raw))


def _words(text: str) -> list[str]:
    return [part for part in text.split() if part]


def _token(text: str, index: int) -> PageToken:
    shift = min(index, 40) * 0.02
    polygon = tuple((x, min(0.98, y + shift)) for x, y in _UNIT)
    return PageToken(
        text=text,
        page=1,
        polygon_source=polygon,
        polygon_norm=polygon,
        engine=ExtractionEngine.VECTOR,
    )


def _skipped_parts(raw: bytes) -> tuple[str, ...]:
    try:
        archive = zipfile.ZipFile(BytesIO(raw))
    except zipfile.BadZipFile:
        return ()
    skipped: list[str] = []
    with archive:
        document_xml = _zip_text(archive, "word/document.xml")
        if "w:txbxContent" in document_xml:
            skipped.append("text_box")
        if "w:sdt" in document_xml:
            skipped.append("content_control")
        for name in archive.namelist():
            lower = name.lower()
            if not lower.startswith("word/header") and not lower.startswith("word/footer"):
                continue
            if "w:t" in _zip_text(archive, name):
                skipped.append("header_footer")
                break
    return tuple(skipped)


def _zip_text(archive: zipfile.ZipFile, name: str) -> str:
    try:
        payload = archive.read(name)
    except KeyError:
        return ""
    return payload.decode("utf-8", errors="ignore")
