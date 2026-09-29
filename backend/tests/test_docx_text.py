"""DOCX: абзацы и таблицы читаются, битый файл не становится UNSUPPORTED."""

from __future__ import annotations

import zipfile
from io import BytesIO

from docx import Document

from kontur.application.docx_text import read_docx_bytes
from kontur.application.process_pipeline import PipelineFile, run_process_pipeline
from kontur.application.scenarios import CompletenessMap
from kontur.domain.models import DocStage
from kontur.domain.statuses import HUMAN_ONLY_STATUSES, Completeness, FindingStatus
from kontur.infrastructure.pdfium_tokens import file_sha256


def _completeness() -> CompletenessMap:
    return {
        DocStage.PD: Completeness.UPLOADED,
        DocStage.RD: Completeness.MISSING,
        DocStage.ID: Completeness.MISSING,
    }


def _docx(paragraph: str, cell: str) -> bytes:
    document = Document()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = cell
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _with_text_box(raw: bytes) -> bytes:
    source = zipfile.ZipFile(BytesIO(raw))
    buffer = BytesIO()
    with source, zipfile.ZipFile(buffer, "w") as target:
        for info in source.infolist():
            payload = source.read(info.filename)
            if info.filename == "word/document.xml":
                payload = payload.replace(b"<w:body>", b"<w:body><w:txbxContent/>", 1)
            target.writestr(info, payload)
    return buffer.getvalue()


def test_paragraph_and_table_become_tokens() -> None:
    parsed = read_docx_bytes(_docx("шифр 12345-PZ", "лист 4"))
    text = " ".join(token.text for token in parsed.tokens)
    assert "12345-PZ" in text
    assert "лист" in text
    assert "4" in text
    assert parsed.skipped == ()


def test_text_box_is_named_and_not_invented() -> None:
    parsed = read_docx_bytes(_with_text_box(_docx("абзац", "ячейка")))
    assert "text_box" in parsed.skipped
    assert "рамка" not in " ".join(token.text for token in parsed.tokens)


def test_pipeline_reads_docx_and_does_not_call_it_unsupported() -> None:
    payload = _docx("шифр 12345-PZ", "лист 4")
    item = PipelineFile(
        file_id="f-docx",
        file_hash=file_sha256(payload),
        filename="note.docx",
        doc_stage=DocStage.PD,
    )
    report = run_process_pipeline(
        object_id="obj-docx",
        completeness=_completeness(),
        files=(item,),
        blobs={"f-docx": payload},
    )
    assert report.pages_built == 1
    text = " ".join(report.parse_errors)
    assert "ACCEPTED_UNPARSED" not in text
    assert "UNSUPPORTED_FORMAT" not in text
    assert all(finding.finding_status not in HUMAN_ONLY_STATUSES for finding in report.findings)
    assert FindingStatus.CONFIRMED_VIOLATION not in {
        finding.finding_status for finding in report.findings
    }


def test_empty_docx_is_an_explicit_miss() -> None:
    payload = _docx("", "")
    parsed = read_docx_bytes(payload)
    assert parsed.tokens == ()
    item = PipelineFile(
        file_id="f-empty",
        file_hash=file_sha256(payload),
        filename="empty.docx",
        doc_stage=DocStage.PD,
    )
    report = run_process_pipeline(
        object_id="obj-empty-docx",
        completeness=_completeness(),
        files=(item,),
        blobs={"f-empty": payload},
    )
    assert report.pages_built == 0
    assert any("DOCX: текст не прочитан" in line for line in report.parse_errors)
    assert all("UNSUPPORTED_FORMAT" not in line for line in report.parse_errors)
