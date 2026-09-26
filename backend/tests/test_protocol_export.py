"""DOCX, XML и PDF протокола из того же провода, что JSON."""

from __future__ import annotations

from io import BytesIO

import pypdfium2 as pdfium
import pytest
from docx import Document
from fastapi.testclient import TestClient

from kontur.application.protocol_export import render_docx, render_pdf, render_xml
from kontur.application.runtime import ProcessWorkspace
from kontur.presentation.api import app

PDF = b"%PDF-1.7\n1 0 obj\n<<>>\nendobj\n"
INSPECTOR = {"Authorization": "Bearer insp-7@obj-1/INSPECTOR"}


def _pdf_text(payload: bytes) -> str:
    document = pdfium.PdfDocument(payload)
    try:
        chunks: list[str] = []
        for index in range(len(document)):
            page = document[index]
            textpage = page.get_textpage()
            try:
                chunks.append(str(textpage.get_text_bounded() or ""))
            finally:
                textpage.close()
                page.close()
        return "\n".join(chunks)
    finally:
        document.close()


def test_docx_and_xml_follow_the_json_wire(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("KONTUR_ALLOW_INSECURE_DEV_AUTH", "true")
    app.state.workspace = ProcessWorkspace()
    client = TestClient(app)
    uploaded = client.post(
        "/api/v1/documents/upload",
        headers=INSPECTOR,
        data={"object_id": "obj-1", "doc_stage": "PD"},
        files=[("files", ("pz.pdf", PDF, "application/pdf"))],
    )
    assert uploaded.status_code == 202
    process_id = uploaded.json()["process_id"]
    docx = client.get(f"/api/v1/processes/{process_id}/protocol.docx", headers=INSPECTOR)
    xml = client.get(f"/api/v1/processes/{process_id}/protocol.xml", headers=INSPECTOR)
    pdf = client.get(f"/api/v1/processes/{process_id}/protocol.pdf", headers=INSPECTOR)
    assert docx.status_code == 200
    assert docx.content.startswith(b"PK")
    assert b"AUTO_NO_DIFFERENCE" not in docx.content
    document = Document(BytesIO(docx.content))
    headings = [paragraph.text for paragraph in document.paragraphs]
    for title in (
        "Статус загрузки документов",
        "Тип проверки",
        "Комплектность",
        "Кандидаты",
        "Подтверждённые нарушения",
        "Нарушение не подтверждено",
        "Подозрения",
        "Требует внимания",
        "Карточки доказательств",
    ):
        assert title in headings
    assert len(document.tables) == 7
    assert xml.status_code == 200
    text = xml.content.decode("utf-8")
    assert "check_type" in text
    assert "upload_status" in text
    assert "evidence_cards" in text
    assert "AUTO_NO_DIFFERENCE" not in text
    assert pdf.status_code == 200
    assert pdf.content.startswith(b"%PDF")
    assert b"AUTO_NO_DIFFERENCE" not in pdf.content
    direct = render_pdf(
        {
            "protocol_id": "p-1",
            "object_id": "obj-1",
            "status": "VERIFICATION_COMPLETED",
            "scenario": "matrix",
            "violation_count": 0,
            "upload_status": {"pd": "PD_UPLOADED", "rd": "RD_MISSING", "id": "ID_MISSING"},
            "sections": {name: [] for name in (
                "completeness",
                "candidates",
                "confirmed",
                "negative_verified",
                "suspicions",
            )},
        }
    )
    assert direct.startswith(b"%PDF")


def test_exports_print_expected_actual_and_fragment() -> None:
    protocol = {
        "protocol_id": "p-1",
        "object_id": "obj-1",
        "status": "VERIFICATION_COMPLETED",
        "scenario": "FULL",
        "violation_count": 0,
        "upload_status": {"pd": "PD_UPLOADED", "rd": "RD_UPLOADED", "id": "ID_MISSING"},
        "sections": {
            "completeness": [],
            "candidates": [
                {
                    "finding_id": "f-1",
                    "rule_code": "PZ-001",
                    "finding_status": "CANDIDATE",
                    "rationale": "расхождение",
                    "expected_value": "1250",
                    "actual_value": "1100",
                    "delta": "150",
                }
            ],
            "confirmed": [],
            "negative_verified": [],
            "suspicions": [],
        },
    }
    cards = {
        "f-1": {
            "file_hash": "ab" * 32,
            "page": "2",
            "polygon_norm": "[[[0.1, 0.2], [0.3, 0.2], [0.3, 0.4]]]",
            "doc_stage": "PD",
            "document_code": "KR-1",
            "revision": "1",
            "approval_basis": "PACKAGE_DEFAULT",
            "tolerance": "abs 1",
        }
    }
    xml = render_xml(protocol, cards).decode("utf-8")
    assert "expected_value" in xml and "1250" in xml
    assert "file_hash" in xml and "polygon_norm" in xml
    document = Document(BytesIO(render_docx(protocol, cards)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Ожидаемое: 1250" in text
    assert "Фактическое: 1100" in text
    assert "Полигон:" in text
    pdf_bytes = render_pdf(protocol, cards)
    assert pdf_bytes.startswith(b"%PDF")
    assert b"AUTO_NO_DIFFERENCE" not in pdf_bytes
    pdf_text = _pdf_text(pdf_bytes)
    for token in ("f-1", "PZ-001", "1250", "1100", "0.1", "ab" * 32):
        assert token in xml
        assert token in text
        assert token in pdf_text
