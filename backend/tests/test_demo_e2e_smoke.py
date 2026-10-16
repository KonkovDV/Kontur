"""Сквозной прогон демо 30.10 через HTTP: учебный комплект → эталон → кандидаты →
решения → протокол → diff редакций → DOCX/PDF/XML.

Тот же путь, что в docs/DEMO_COLD_START.md и у кнопки учебного комплекта
(`POST /demo/kit`, `install_demo_kit`). Тест не пропускается: без TTF с
кириллицей он падает, а не молчит (шрифт лежит в vendor/fonts).
"""

from __future__ import annotations

from io import BytesIO
from time import perf_counter

import pypdfium2 as pdfium
import pytest
from docx import Document
from fastapi.testclient import TestClient
from lxml import etree
from pdf_fixtures import contest_slice_font

from kontur.application.runtime import ProcessWorkspace
from kontur.cli.run_package import _schema_validator
from kontur.evaluation.agent_dumps import repo_root
from kontur.evaluation.demo_cold_start import DEMO_CODES
from kontur.evaluation.demo_kit import DRAFT_FILE_ID, ETALON_FILE_ID
from kontur.presentation.api import app

INSPECTOR_ID = "insp-demo"
HEADERS = {"Authorization": f"Bearer {INSPECTOR_ID}@OBJ-DEMO-E2E/INSPECTOR"}
#: Бюджет всего сценария в CI. Порог демо по ТЗ — 30 минут на комплект;
#: здесь синтетика из четырёх листов: локально около 1-2 с, запас на медленный раннер CI.
BUDGET_SECONDS = 20.0

if contest_slice_font() is None:
    pytest.fail("нет TTF с кириллицей: vendor/fonts/NotoSans-Regular.ttf пропал", pytrace=False)


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("KONTUR_ALLOW_INSECURE_DEV_AUTH", "true")
    app.state.workspace = ProcessWorkspace()
    return TestClient(app)


def _findings(client: TestClient, process_id: str) -> dict[str, dict[str, object]]:
    response = client.get(f"/api/v1/processes/{process_id}/findings", headers=HEADERS)
    assert response.status_code == 200
    return {row["rule_code"]: row for row in response.json()["findings"]}


def _protocol(
    client: TestClient, process_id: str, *, validate: bool = True
) -> dict[str, object]:
    response = client.get(f"/api/v1/processes/{process_id}/protocol", headers=HEADERS)
    assert response.status_code == 200
    body: dict[str, object] = response.json()
    if validate:
        _schema_validator("protocol.schema.json").validate(body)
    return body


def _post(client: TestClient, path: str, body: dict[str, object]) -> dict[str, object]:
    response = client.post(f"/api/v1{path}", headers=HEADERS, json=body)
    assert response.status_code == 200, (path, response.status_code, response.text)
    payload: dict[str, object] = response.json()
    return payload


def _run_to_finalize(client: TestClient, timings: dict[str, float]) -> str:
    # 1. Учебный комплект: две ПД (черновик и утверждённая), РД, ИД. Эталона нет.
    tick = perf_counter()
    kit = _post(client, "/demo/kit", {})
    timings["kit"] = perf_counter() - tick
    process_id = str(kit["process_id"])
    assert kit["etalon_file_id"] == ETALON_FILE_ID
    before = _findings(client, process_id)
    assert {before[code]["finding_status"] for code in DEMO_CODES} == {"CLARIFICATION_REQUIRED"}
    assert "CONFIRMED_VIOLATION" not in {row["finding_status"] for row in before.values()}

    # 2. Инспектор назначает эталон: ровно три кандидата.
    tick = perf_counter()
    _post(
        client,
        f"/processes/{process_id}/revisions/{ETALON_FILE_ID}/select",
        {"inspector_id": INSPECTOR_ID, "comment": "Утверждённая редакция тома."},
    )
    timings["select"] = perf_counter() - tick
    after = _findings(client, process_id)
    candidates = {code for code, row in after.items() if row["finding_status"] == "CANDIDATE"}
    assert candidates == set(DEMO_CODES)
    for code in DEMO_CODES:
        row = after[code]
        assert row["places"], code
        assert row["evidence_group_id"], code

    ready = _protocol(client, process_id)
    assert ready["status"] == "READY"
    assert ready["violation_count"] == 0

    # 3. Diff редакций: пара черновик → эталон в комплекте не зафиксирована,
    #    порядок не угадывается: у эталона нет предшественника (409). Черновик после
    #    выбора эталона не попадает в группы доказательств, поэтому diff его не видит (404).
    diff_url = f"/api/v1/processes/{process_id}/revisions/{{}}/diff"
    assert client.get(diff_url.format(ETALON_FILE_ID), headers=HEADERS).status_code == 409
    assert client.get(diff_url.format(DRAFT_FILE_ID), headers=HEADERS).status_code == 404
    assert "revision_changes" not in ready

    # 4. Решения: два подтверждения, один отказ с причиной.
    tick = perf_counter()
    decisions = (
        ("PZ-001", "CONFIRM", None),
        ("KR-055", "CONFIRM", None),
        ("AR-041", "REJECT", "APPROVED_CHANGE_EXISTS"),
    )
    for code, action, reason in decisions:
        body: dict[str, object] = {
            "action": action,
            "inspector_id": INSPECTOR_ID,
            "comment": f"Демо: {code}.",
        }
        if reason is not None:
            body["reason_code"] = reason
        _post(client, f"/findings/{after[code]['finding_id']}/review", body)
    _post(client, f"/processes/{process_id}/complete", {"inspector_id": INSPECTOR_ID})
    _post(client, f"/processes/{process_id}/finalize", {"inspector_id": INSPECTOR_ID})
    timings["review_finalize"] = perf_counter() - tick
    return process_id


def test_demo_scenario_end_to_end(client: TestClient) -> None:
    timings: dict[str, float] = {}
    started = perf_counter()
    process_id = _run_to_finalize(client, timings)

    final = _protocol(client, process_id)
    assert final["status"] == "PROTOCOL_FINALIZED"
    assert final["violation_count"] == 2
    sections = final["sections"]
    assert isinstance(sections, dict)
    assert {row["rule_code"] for row in sections["confirmed"]} == {"PZ-001", "KR-055"}
    assert {row["rule_code"] for row in sections["negative_verified"]} == {"AR-041"}
    assert "AUTO_NO_DIFFERENCE" not in str(final)

    # 5. Экспорт: файлы не пустые и разбираются.
    tick = perf_counter()
    base = f"/api/v1/processes/{process_id}/protocol"
    docx = client.get(f"{base}.docx", headers=HEADERS)
    assert docx.status_code == 200 and docx.content
    text = "\n".join(p.text for p in Document(BytesIO(docx.content)).paragraphs)
    assert "PZ-001" in text and "KR-055" in text
    pdf = client.get(f"{base}.pdf", headers=HEADERS)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    document = pdfium.PdfDocument(pdf.content)
    try:
        assert len(document) >= 1
        page_text = "".join(
            document[index].get_textpage().get_text_range() for index in range(len(document))
        )
    finally:
        document.close()
    assert "PZ-001" in page_text
    xml = client.get(f"{base}.xml", headers=HEADERS)
    assert xml.status_code == 200 and xml.content
    schema = etree.XMLSchema(
        etree.parse(str(repo_root() / "contracts" / "schemas" / "protocol-export.xsd"))
    )
    root = etree.fromstring(xml.content)
    assert schema.validate(root)
    assert root.findtext("violation_count") == "2"
    timings["export"] = perf_counter() - tick

    total = perf_counter() - started
    print("demo e2e timings:", {key: round(value, 2) for key, value in timings.items()})
    print(f"demo e2e total: {total:.2f}s (budget {BUDGET_SECONDS:.0f}s)")
    assert total < BUDGET_SECONDS



def test_finalized_protocol_and_exports_match_contracts(client: TestClient) -> None:
    process_id = _run_to_finalize(client, {})
    final = _protocol(client, process_id)
    sections = final["sections"]
    assert isinstance(sections, dict)
    assert sections["completeness"], "таблица комплектности в финальном протоколе"
    assert final["violation_count"] == 2
    base = f"/api/v1/processes/{process_id}/protocol"
    xml = client.get(f"{base}.xml", headers=HEADERS)
    schema = etree.XMLSchema(
        etree.parse(str(repo_root() / "contracts" / "schemas" / "protocol-export.xsd"))
    )
    assert schema.validate(etree.fromstring(xml.content))
    docx = client.get(f"{base}.docx", headers=HEADERS)
    assert Document(BytesIO(docx.content)).paragraphs
    pdf = client.get(f"{base}.pdf", headers=HEADERS)
    assert pdf.content.startswith(b"%PDF")


def test_outbox_still_carries_only_the_rin_body(client: TestClient) -> None:
    process_id = _run_to_finalize(client, {})
    workspace = app.state.workspace
    record = workspace.get(process_id)
    assert record is not None and record.protocol_id is not None
    wire = workspace.load_protocol(record.protocol_id)
    assert wire is not None
    assert set(wire["sections"]) == {"confirmed", "negative_verified"}
